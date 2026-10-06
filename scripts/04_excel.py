"""Schritt 4: Excel-Tabelle (Übersicht mit Formeln + Blatt mit allen Flächen).

Liest  output/<dateinamen.gpkg> (Layer flaechen, ohne Geometrie)
Schreibt output/<dateinamen.excel>

Die Übersicht rechnet mit SUMIFS/COUNTIF auf dem Blatt «Flächen». Excel und Numbers rechnen beim
Öffnen neu. Damit auch Vorschauen ohne Rechenwerk (Quick Look, Browser) Zahlen zeigen, schreibt das
Skript die in Python berechneten Ergebnisse zusätzlich als zwischengespeicherte Werte in die Datei.
"""
import re
import zipfile

import pyogrio
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from common import CFG, FARMS, GROUPS, kantone_text, out_path, teile

d = pyogrio.read_dataframe(out_path("gpkg"), layer="flaechen", read_geometry=False)
d = d.sort_values(["betrieb", "flaeche_a"], ascending=[True, False])
farm_names = [f["name"] for f in FARMS]
kts = [k for k in CFG["kantone"] if (d.betriebsnummer.str[:2] == k).any()]


def F(**k):
    return Font(name="Arial", **k)


wb = Workbook()
ov = wb.active
ov.title = "Übersicht"
fl = wb.create_sheet("Flächen")
hdr = PatternFill("solid", start_color="2C5B45")

cols = [("Betrieb", "betrieb", 12), ("Betriebsnummer", "betriebsnummer", 16), ("Kultur", "kultur", 48),
        ("LNF-Code", "lnf_code", 9), ("Kulturgruppe", "kulturgruppe", 24), ("Fläche (a)", "flaeche_a", 11),
        ("Fläche (ha)", "flaeche_ha", 11), ("Gemeinde", "gemeinde", 20), ("Kanton", "kanton", 7),
        ("Programme", "programme", 60), ("beitragsberechtigt", "beitragsberechtigt", 12), ("Bezugsjahr", "bezugsjahr", 10),
        ("Teilbetrieb", "teilbetrieb", 24), ("Bio-Status", "bio_status", 38)]
for j, (h, _, wdt) in enumerate(cols, 1):
    c = fl.cell(1, j, h)
    c.font = F(bold=True, color="FFFFFF")
    c.fill = hdr
    fl.column_dimensions[get_column_letter(j)].width = wdt
for i, r in enumerate(d.itertuples(index=False), 2):
    for j, (_, k, _) in enumerate(cols, 1):
        v = getattr(r, k)
        if k == "programme":
            v = (v or "").replace(";", "; ").replace("Kein Programm", "").strip("; ")
        if k == "beitragsberechtigt":
            v = "ja" if v else "nein"
        if k in ("lnf_code", "bezugsjahr"):
            v = int(v)
        c = fl.cell(i, j, v)
        c.font = F()
        if k == "flaeche_a":
            c.number_format = "#,##0.0"
        if k == "flaeche_ha":
            c.number_format = "#,##0.000"
n = len(d) + 1
fl.freeze_panes = "A2"
fl.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{n}"

ov["A1"] = f"{CFG['titel']} – Fläche nach Kulturgruppe (ha)"
ov["A1"].font = F(bold=True, size=13)
ov["A2"] = f"Quelle: {kantone_text(kts)}, Nutzungsflächen und Bewirtschaftungseinheiten (geodienste.ch), {CFG['stand']}. Berechnet aus Blatt «Flächen»."
ov["A2"].font = F(italic=True, size=9, color="5D6A62")
nr_text = " · ".join(f"{f['name']} {' + '.join(f['nrs'])}" + (f" ({f['note']})" if f["note"] else "") for f in FARMS)
ov["A3"] = "Betriebsnummern: " + nr_text
ov["A3"].font = F(italic=True, size=9, color="5D6A62")

hr = 5
tj = len(farm_names) + 2
ov.cell(hr, 1, "Kulturgruppe")
for j, f in enumerate(farm_names, 2):
    ov.cell(hr, j, f)
ov.cell(hr, tj, "Total")
for j in range(1, tj + 1):
    c = ov.cell(hr, j)
    c.font = F(bold=True, color="FFFFFF")
    c.fill = hdr
fmt = '#,##0.0;-#,##0.0;"–"'
cache = {}   # Zelle -> in Python berechneter Wert der Formel (für Vorschauen)
ha_ = lambda m: round(float(d[m].flaeche_ha.sum()), 6)
for i, g in enumerate(GROUPS, hr + 1):
    ov.cell(i, 1, g).font = F()
    for j in range(2, tj):
        L = get_column_letter(j)
        c = ov.cell(i, j, f"=SUMIFS(Flächen!$G$2:$G${n},Flächen!$A$2:$A${n},{L}${hr},Flächen!$E$2:$E${n},$A{i})")
        cache[f"{L}{i}"] = ha_((d.betrieb == farm_names[j - 2]) & (d.kulturgruppe == g))
        c.number_format = fmt
        c.font = F()
    c = ov.cell(i, tj, f"=SUM(B{i}:{get_column_letter(tj - 1)}{i})")
    cache[f"{get_column_letter(tj)}{i}"] = ha_(d.kulturgruppe == g)
    c.number_format = fmt
    c.font = F(bold=True)
tr = hr + len(GROUPS) + 1
ov.cell(tr, 1, "Total").font = F(bold=True)
for j in range(2, tj + 1):
    L = get_column_letter(j)
    c = ov.cell(tr, j, f"=SUM({L}{hr + 1}:{L}{tr - 1})")
    cache[f"{L}{tr}"] = ha_(d.betrieb == farm_names[j - 2]) if j < tj else ha_(d.betrieb.notna())
    c.number_format = "#,##0.0"
    c.font = F(bold=True)
    c.border = Border(top=Side(style="thin"))
ar = tr + 1
ov.cell(ar, 1, "Anzahl Flächen").font = F()
for j in range(2, tj):
    ov.cell(ar, j, f"=COUNTIF(Flächen!$A$2:$A${n},{get_column_letter(j)}${hr})").font = F()
    cache[f"{get_column_letter(j)}{ar}"] = int((d.betrieb == farm_names[j - 2]).sum())
c = ov.cell(ar, tj, f"=SUM(B{ar}:{get_column_letter(tj - 1)}{ar})")
cache[f"{get_column_letter(tj)}{ar}"] = len(d)
c.font = F(bold=True)
# Teilbetriebe (Betriebe mit mehreren Firmen/Standorten, getrennt über Betriebsnummer bzw. Produktionsstätte)
mehr = [(f, t) for f in FARMS if f.get("teile") for t in teile(f)]
if mehr:
    tr0 = ar + 2
    ov.cell(tr0, 1, "Teilbetriebe (aufgeteilt nach Betriebsnummer bzw. amtlicher Produktionsstätte)").font = F(bold=True, size=11)
    kopf = ["Teilbetrieb", "Betrieb", "Fläche (ha)", "Gemüse (ha)", "Anzahl Flächen", "Bio-Status"]
    for j, h in enumerate(kopf, 1):
        c = ov.cell(tr0 + 1, j, h)
        c.font = F(bold=True, color="FFFFFF")
        c.fill = hdr
    for i, (f, t) in enumerate(mehr, tr0 + 2):
        ov.cell(i, 1, t["name"]).font = F()
        ov.cell(i, 2, f["name"]).font = F()
        c = ov.cell(i, 3, f"=SUMIFS(Flächen!$G$2:$G${n},Flächen!$M$2:$M${n},$A{i})")
        c.number_format, c.font = fmt, F()
        cache[f"C{i}"] = ha_(d.teilbetrieb == t["name"])
        c = ov.cell(i, 4, f'=SUMIFS(Flächen!$G$2:$G${n},Flächen!$M$2:$M${n},$A{i},Flächen!$E$2:$E${n},"Freilandgemüse")'
                          f'+SUMIFS(Flächen!$G$2:$G${n},Flächen!$M$2:$M${n},$A{i},Flächen!$E$2:$E${n},"Gewächshaus / geschützt")')
        c.number_format, c.font = fmt, F()
        cache[f"D{i}"] = ha_((d.teilbetrieb == t["name"]) & d.kulturgruppe.isin(["Freilandgemüse", "Gewächshaus / geschützt"]))
        ov.cell(i, 5, f"=COUNTIF(Flächen!$M$2:$M${n},$A{i})").font = F()
        cache[f"E{i}"] = int((d.teilbetrieb == t["name"]).sum())
        ov.cell(i, 6, t.get("bio_text", "")).font = F(size=9)
ov.column_dimensions["A"].width = 28
for j in range(2, tj + 1):
    ov.column_dimensions[get_column_letter(j)].width = 13
ov.freeze_panes = f"B{hr + 1}"

p = out_path("excel")
wb.save(p)


def werte_einsetzen(pfad, blatt_xml, werte):
    """Zwischengespeicherte Ergebnisse in die Formelzellen schreiben (<v>…</v>); Excel rechnet beim Öffnen trotzdem neu."""
    tmp = pfad.with_suffix(".tmp")
    with zipfile.ZipFile(pfad) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            daten = zin.read(it.filename)
            if it.filename == blatt_xml:
                x = daten.decode("utf-8")
                x = re.sub(r'(<c r="([A-Z]+[0-9]+)"[^>]*><f>[^<]*</f>)<v ?/>',
                           lambda m: f"{m.group(1)}<v>{werte[m.group(2)]}</v>" if m.group(2) in werte else m.group(0), x)
                daten = x.encode("utf-8")
            zout.writestr(it, daten)
    tmp.replace(pfad)


werte_einsetzen(p, "xl/worksheets/sheet1.xml", cache)   # Blatt «Übersicht» ist das erste
print(f"{len(d)} Zeilen, {round(d.flaeche_ha.sum(), 2)} ha -> output/{p.name}")
